import os
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
from pathlib import Path

# Color mapping for each class_type
COLOR_MAP = {
    "SAM": "FF9999",  # Soft Red fill
    "MAM": "FFFF99",  # Soft Yellow fill
    "ONC": "99FF99"   # Soft Green fill
}

def get_image_identifiers(image_folder):
    """
    Extracts filename identifiers (both full name and stem without extension)
    from the specified image folder. Handles both files and subdirectories.
    """
    if not image_folder or not os.path.exists(image_folder):
        print(f"Warning: Image folder does not exist or not provided: '{image_folder}'")
        return set()
    
    image_names = set()
    
    for entry in os.listdir(image_folder):
        entry_clean = entry.strip().lower()
        entry_stem = Path(entry).stem.strip().lower()
        image_names.add(entry_clean)
        image_names.add(entry_stem)
                
    return image_names


def highlight_multi_folders(
    input_dataset_path: str,
    folder_mapping: dict,
    output_excel_path: str = None,
    output_csv_path: str = None,
    filename_col: str = "File_Name",
    class_type_col: str = "Class_Type",
    match_class_type: bool = True
):
    """
    Highlights dataset rows for SAM (Red), MAM (Yellow), and ONC (Green) in ONE GO.

    Parameters:
    -----------
    input_dataset_path : str
        Path to input CSV or Excel dataset (.csv, .xlsx)
    folder_mapping : dict
        Dictionary mapping class_type ('SAM', 'MAM', 'ONC') to image folder paths.
    output_excel_path : str, optional
        Destination path for saved highlighted .xlsx file.
    output_csv_path : str, optional
        Destination path for saved highlighted .csv file.
    filename_col : str, optional
        Column name in dataset containing image filenames (default: 'File_Name').
    class_type_col : str, optional
        Column name in dataset containing class category (default: 'Class_Type').
    match_class_type : bool, optional
        If True (default), ensures dataset row's Class_Type matches folder class ('sam', 'mam', 'onc').
    """
    if output_excel_path is None:
        base_name = os.path.splitext(input_dataset_path)[0]
        output_excel_path = f"{base_name}_highlighted.xlsx"
    elif not output_excel_path.endswith('.xlsx'):
        output_excel_path += '.xlsx'

    # Pre-extract image identifiers and create fill styles for each class
    class_data = {}
    for c_type, folder_path in folder_mapping.items():
        c_key = c_type.strip().upper()
        if c_key not in COLOR_MAP:
            print(f"Warning: Unknown class_type '{c_type}'. Supported: {list(COLOR_MAP.keys())}")
            continue
        
        hex_color = COLOR_MAP[c_key]
        fill_style = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        
        print(f"--> Extracting image names for [{c_key}] (Color: #{hex_color}) from:\n    {folder_path}")
        image_ids = get_image_identifiers(folder_path)
        print(f"    Found {len(image_ids)} unique image identifier keys for [{c_key}].")
        
        class_data[c_key] = {
            "fill": fill_style,
            "image_ids": image_ids,
            "match_count": 0,
            "color_name": "Red" if c_key == "SAM" else "Yellow" if c_key == "MAM" else "Green"
        }

    # Load dataset into pandas first
    print(f"\n--> Reading input dataset: {input_dataset_path}")
    if input_dataset_path.endswith('.xlsx') or input_dataset_path.endswith('.xls'):
        df = pd.read_excel(input_dataset_path)
    else:
        df = pd.read_csv(input_dataset_path)

    if filename_col not in df.columns:
        raise KeyError(f"Column '{filename_col}' not found in dataset columns: {list(df.columns[:5])}...")

    # Add Highlighted columns for CSV export
    df['Highlighted_Class'] = 'None'
    df['Highlighted_Color'] = 'None'

    # Load openpyxl Workbook for cell coloring
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Highlighted Data"
    
    header_cols = list(df.columns)
    ws.append(header_cols)

    total_rows = len(df)
    print(f"--> Scanning {total_rows} dataset rows and applying highlights (match_class_type={match_class_type})...")

    for r_idx, row in df.iterrows():
        row_num = r_idx + 2  # 1-based index + header row
        cell_val = str(row[filename_col] or '').strip()
        row_class = str(row.get(class_type_col, '') or '').strip().upper()

        cell_val_lower = cell_val.lower()
        cell_stem_lower = Path(cell_val).stem.lower()

        matched_class = None

        # Check match against all classes
        for c_key, c_info in class_data.items():
            # If match_class_type is enabled, skip if CSV Class_Type doesn't match target class
            if match_class_type and row_class != c_key:
                continue

            if cell_val_lower in c_info["image_ids"] or cell_stem_lower in c_info["image_ids"]:
                c_info["match_count"] += 1
                matched_class = c_key
                df.at[r_idx, 'Highlighted_Class'] = matched_class
                df.at[r_idx, 'Highlighted_Color'] = c_info["color_name"]
                break

        # Append row to Excel worksheet
        row_values = [row[col] for col in header_cols]
        ws.append(row_values)

        # Fill color if matched
        if matched_class and matched_class in class_data:
            fill_style = class_data[matched_class]["fill"]
            for c_col in range(1, ws.max_column + 1):
                ws.cell(row=row_num, column=c_col).fill = fill_style

    # Print Summary
    print("\n" + "=" * 60)
    print("HIGHLIGHTING SUMMARY (ONE GO):")
    print("=" * 60)
    for c_key, c_info in class_data.items():
        print(f" - {c_key} ({c_info['color_name']} / #{COLOR_MAP[c_key]}): Highlighted {c_info['match_count']} rows")
    print("=" * 60)

    # Save highlighted Excel file
    wb.save(output_excel_path)
    print(f"\n--> Successfully saved highlighted Excel file (.xlsx) to:\n    {output_excel_path}")

    # Save highlighted CSV file
    if output_csv_path is None:
        base_name = os.path.splitext(input_dataset_path)[0]
        output_csv_path = f"{base_name}_highlighted.csv"
    
    df.to_csv(output_csv_path, index=False)
    print(f"--> Successfully saved highlighted CSV file (.csv) to:\n    {output_csv_path}\n")

    return output_excel_path, output_csv_path


def highlight_dataset_by_class(
    input_dataset_path: str,
    image_folder_path: str,
    class_type: str,
    output_excel_path: str = None,
    output_csv_path: str = None,
    filename_col: str = "File_Name",
    match_class_type: bool = True
):
    """
    Highlights dataset rows for a single class_type (SAM, MAM, or ONC).
    """
    return highlight_multi_folders(
        input_dataset_path=input_dataset_path,
        folder_mapping={class_type: image_folder_path},
        output_excel_path=output_excel_path,
        output_csv_path=output_csv_path,
        filename_col=filename_col,
        match_class_type=match_class_type
    )


if __name__ == "__main__":
    # Execution with actual folders
    dataset_csv = r"GM_out_of_sample_classification-main\GM_out_of_sample_classification-ap-govt\final_dataset_centred.csv"

    folder_mapping = {
        "SAM": r"C:\complete_global_child_arm_dataset\SAM\sam_filtered_image\ok_images",
        "MAM": r"C:\complete_global_child_arm_dataset\MAM\mam_filtered_image\ok",
        "ONC": r"C:\noc_filtered_images\ok images"
    }

    output_excel = r"GM_out_of_sample_classification-main\GM_out_of_sample_classification-ap-govt\final_dataset_highlighted.xlsx"
    output_csv = r"GM_out_of_sample_classification-main\GM_out_of_sample_classification-ap-govt\final_dataset_highlighted.csv"

    if os.path.exists(dataset_csv):
        highlight_multi_folders(
            input_dataset_path=dataset_csv,
            folder_mapping=folder_mapping,
            output_excel_path=output_excel,
            output_csv_path=output_csv,
            match_class_type=True
        )
